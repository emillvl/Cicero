"""Offline integration tests. No keys, model inference, or network needed.
Run: python -m pytest -q (or python -m unittest -v test_translator)
"""
import contextlib
import io
import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import patch

import interactive_translator as app


SOURCE = '# Harbor\n\nMira arrived on Tuesday carrying 17 blue books. She did not open the red box.\n\n- Keep the books dry.'
DRAFT = '# Puerto\n\nMira llegó el martes con 17 libros azules. No abrió la caja roja.\n\n- Mantén secos los libros.'
FIXED = DRAFT.replace('Mantén secos', 'Mantenga secos')
PASS = 'VERDICT: PASS'
FIX = 'VERDICT: FIX\nUse a consistent formal voice in the list item.'
SECRET = 'FAKE_PRIVATE_CREDENTIAL_FOR_OFFLINE_TEST_ONLY'


class FakeModels:
    def __init__(self, responses):
        self.calls = []
        # A dict maps a role to its queued replies; a list is consumed in call order.
        # Role-keyed replies keep tests deterministic now that the two reviewers run concurrently.
        if isinstance(responses, dict):
            self._by_role = {role: iter(value if isinstance(value, list) else [value])
                             for role, value in responses.items()}
            self._sequence = None
        else:
            self._by_role = None
            self._sequence = iter(responses)

    def call(self, cfg, role, system, payload, **kwargs):
        self.calls.append((role, payload))
        result = next(self._by_role[role]) if self._by_role is not None else next(self._sequence)
        if isinstance(result, BaseException):
            raise result
        return result


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.source = self.folder / 'source.txt'
        self.source.write_text(SOURCE, encoding='utf-8')
        self.original = self.source.read_bytes()
        self.settings = app.Settings(language='Spanish')
        cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], SECRET, 128000, 4096)
        self.configs = dict.fromkeys(app.ROLES, cfg)
        self.chunks = [app.text_blocks(SOURCE)]
        self.identity = app.identity_for(self.source, self.chunks, self.configs, self.settings)
        self.cp = app.Checkpoint(self.folder, self.identity)

    def tearDown(self):
        self.temp.cleanup()

    def run_pipeline(self, responses, cp=None):
        fake = FakeModels(responses)
        pipeline = app.Pipeline(fake, self.configs, self.settings, self.chunks, cp or self.cp)
        with contextlib.redirect_stdout(io.StringIO()):
            result = pipeline.run()
        return result, fake

    def test_balanced_plain_text_and_complete_resume_no_calls(self):
        text, fake = self.run_pipeline([DRAFT, PASS, PASS])
        self.assertEqual(text, DRAFT + '\n')
        # The two reviewers run concurrently, so only their set matters.
        self.assertEqual(Counter(x[0] for x in fake.calls), Counter(translator=1, fluency=1, accuracy=1))
        self.assertEqual(self.source.read_bytes(), self.original)
        cp = app.Checkpoint(self.folder, self.identity)
        text, fake = self.run_pipeline([], cp)
        self.assertEqual(len(fake.calls), 0)
        self.assertTrue(cp.data['complete'])
        for f in self.folder.iterdir():
            self.assertNotIn(SECRET.encode(), f.read_bytes())

    def test_corrections_are_audited_and_notes_not_delivered(self):
        text, fake = self.run_pipeline([DRAFT, FIX, PASS, FIXED, PASS])
        self.assertEqual(text, FIXED + '\n')
        self.assertEqual([x[0] for x in fake.calls][-2:], ['fixer', 'auditor'])
        self.assertNotIn('VERDICT:', (self.folder / 'translation.txt').read_text(encoding='utf-8'))

    def test_full_always_audits(self):
        self.settings.thoroughness = 'full'
        _, fake = self.run_pipeline([DRAFT, PASS, PASS, PASS])
        self.assertEqual(fake.calls[-1][0], 'auditor')

    def test_failed_review_is_not_approval_and_resume_skips_translation(self):
        with self.assertRaises(app.TranslationError):
            self.run_pipeline({'translator': DRAFT, 'fluency': 'looks good', 'accuracy': PASS})
        self.assertEqual(self.cp.data['pending'], '0:fluency')
        # The independent accuracy review still completed while fluency failed.
        self.assertEqual(set(self.cp.data['stages']), {'0:translation', '0:accuracy'})
        self.assertFalse((self.folder / 'translation.txt').exists())
        partial = (self.folder / 'translation.partial.txt').read_text(encoding='utf-8')
        self.assertIn(DRAFT, partial)
        self.assertIn('review incomplete', partial)
        cp = app.Checkpoint(self.folder, self.identity)
        with self.assertRaises(app.TranslationError):
            self.run_pipeline([], cp)
        cp.data['pending'] = None  # Represents explicit retry consent in the CLI.
        cp.save()
        _, fake = self.run_pipeline({'fluency': PASS}, cp)
        self.assertEqual([x[0] for x in fake.calls], ['fluency'])

    def test_timeout_preserves_previous_stages_and_source(self):
        error = app.APIError(0, 'ReadTimeout. No automatic retry.')
        with self.assertRaises(app.APIError):
            self.run_pipeline([DRAFT, PASS, error])
        cp = app.Checkpoint(self.folder, self.identity)
        self.assertEqual(set(cp.data['stages']), {'0:translation', '0:fluency'})
        self.assertEqual(cp.data['pending'], '0:accuracy')
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_interrupted_fix_resumes_without_paid_review_replay(self):
        with self.assertRaises(app.APIError):
            self.run_pipeline([DRAFT, FIX, PASS, FIXED, app.APIError(0, 'read timeout')])
        cp = app.Checkpoint(self.folder, self.identity)
        cp.data['pending'] = None
        cp.save()
        result, fake = self.run_pipeline([PASS], cp)
        self.assertEqual(result, FIXED + '\n')
        self.assertEqual([x[0] for x in fake.calls], ['auditor'])

    def test_two_failed_audits_stop_and_do_not_repeat_on_restart(self):
        with self.assertRaises(app.TranslationError):
            self.run_pipeline([DRAFT, FIX, PASS, FIXED, FIX, FIXED, FIX])
        self.assertFalse(self.cp.data['complete'])
        self.assertIsNone(self.cp.data['pending'])
        with self.assertRaises(app.TranslationError):
            self.run_pipeline([])
        self.assertFalse((self.folder / 'translation.txt').exists())

    def test_mismatch_never_overwrites_checkpoint(self):
        original = self.cp.path.read_bytes()
        for field, value in [('source_sha256', 'changed'), ('version', 'new'), ('settings', {})]:
            bad = dict(self.identity, **{field: value})
            with self.assertRaises(app.Mismatch):
                app.Checkpoint(self.folder, bad)
            self.assertEqual(original, self.cp.path.read_bytes())

    def test_corrupt_checkpoint_not_overwritten(self):
        self.cp.path.write_text('{bad', encoding='utf-8')
        with self.assertRaises(app.TranslationError):
            app.Checkpoint(self.folder, self.identity)
        self.assertEqual(self.cp.path.read_text(), '{bad')

    def test_atomic_write_failure_keeps_last_checkpoint(self):
        original = self.cp.path.read_bytes()
        with patch.object(app.os, 'replace', side_effect=OSError('simulated disk error')):
            with self.assertRaises(OSError):
                self.cp.save()
        self.assertEqual(self.cp.path.read_bytes(), original)
        self.assertFalse(list(self.folder.glob('.checkpoint*')))

    def test_export_failure_keeps_utf8_txt(self):
        text, _ = self.run_pipeline([DRAFT, PASS, PASS])
        with patch.object(app, 'export_docx', side_effect=RuntimeError('export unavailable')):
            with contextlib.redirect_stdout(io.StringIO()) as messages:
                app.rich_exports(self.folder, text, ('docx',))
        self.assertIn('UTF-8 translation.txt remains', messages.getvalue())
        self.assertEqual((self.folder / 'translation.txt').read_text(encoding='utf-8'), text)
        self.assertFalse((self.folder / 'translation.docx').exists())

    def test_secret_in_response_never_persisted(self):
        with self.assertRaises(app.TranslationError):
            self.run_pipeline([DRAFT + SECRET])
        self.assertNotIn(SECRET, self.cp.path.read_text())

    def test_multiple_passages_context_does_not_duplicate(self):
        self.chunks = [[app.Block('First passage')], [app.Block('Second passage')]]
        text, fake = self.run_pipeline(['Primer pasaje', PASS, PASS, 'Segundo pasaje', PASS, PASS])
        self.assertEqual(text, 'Primer pasaje\n\nSegundo pasaje\n')
        self.assertIn('Primer pasaje', fake.calls[3][1])

    def test_translation_can_merge_paragraphs_and_drop_formatting(self):
        plain = DRAFT.replace('# ', '').replace('- ', '').replace('\n\n', ' ')
        text, fake = self.run_pipeline([plain, PASS, PASS])
        self.assertEqual(text, plain + '\n')
        self.assertEqual(Counter(x[0] for x in fake.calls), Counter(translator=1, fluency=1, accuracy=1))
        self.assertIn(plain, (self.folder / 'passage-1-translation.received.txt').read_text(encoding='utf-8'))
        self.assertNotIn('# Harbor', fake.calls[0][1])

    def test_translation_can_add_paragraph_breaks(self):
        draft = DRAFT.replace('. No abrió', '.\n\nNo abrió')
        text, _ = self.run_pipeline([draft, PASS, PASS])
        self.assertEqual(text, draft + '\n')

    def test_reviewer_catches_missing_content_without_format_matching(self):
        draft = DRAFT.replace(' No abrió la caja roja.', '').replace('\n\n', '\n')
        text, fake = self.run_pipeline([draft, PASS, 'VERDICT: FIX\nThe red-box sentence is missing.', DRAFT, PASS])
        self.assertEqual(text, DRAFT + '\n')
        self.assertEqual(fake.calls[-2][0], 'fixer')

    def test_rejected_translation_is_saved_as_unreviewed(self):
        invalid = 'VERDICT: PASS\n' + DRAFT
        with self.assertRaises(app.TranslationError):
            self.run_pipeline([invalid])
        received = (self.folder / 'passage-1-translation.received.txt').read_text(encoding='utf-8')
        self.assertIn('UNREVIEWED', received)
        self.assertIn(invalid, received)
        self.assertNotIn('0:translation', self.cp.data['stages'])
        self.assertFalse((self.folder / 'translation.txt').exists())

    def test_known_v1_checkpoint_upgrades_and_preserves_paid_stages(self):
        old = dict(self.identity, version='cicero-plain-1',
                   prompt_sha256=app.legacy_prompt_hash(self.settings.language))
        self.cp.data['identity'] = old
        self.cp.data['stages'] = {'0:translation': {'text': DRAFT, 'seconds': 42}}
        self.cp.data['pending'] = '0:fluency'
        self.cp.save()
        original = self.cp.path.read_text(encoding='utf-8')
        with contextlib.redirect_stdout(io.StringIO()):
            cp = app.Checkpoint(self.folder, self.identity)
        self.assertEqual(cp.data['identity'], self.identity)
        self.assertEqual(cp.data['pending'], '0:fluency')
        self.assertEqual(cp.data['identity_history'], [old])
        self.assertEqual((self.folder / 'checkpoint.before-text-only.json').read_text(encoding='utf-8'), original)
        cp.data['pending'] = None
        cp.save()
        _, fake = self.run_pipeline([PASS, PASS], cp)
        self.assertEqual([x[0] for x in fake.calls], ['fluency', 'accuracy'])

    def test_v1_upgrade_refuses_other_configuration_changes(self):
        previous = dict(self.identity, version='cicero-plain-1',
                        prompt_sha256=app.legacy_prompt_hash(self.settings.language))
        for field, value in [('source_sha256', 'changed'), ('roles', {}),
                             ('chunks_sha256', 'changed'), ('prompt_sha256', 'unknown')]:
            with self.subTest(field=field):
                self.cp.data['identity'] = dict(previous, **{field: value})
                self.cp.save()
                before = self.cp.path.read_bytes()
                with self.assertRaises(app.Mismatch):
                    app.Checkpoint(self.folder, self.identity)
                self.assertEqual(self.cp.path.read_bytes(), before)

    def test_v2_upgrade_resumes_at_fluency_without_retranslation(self):
        old = dict(self.identity, version='cicero-plain-2',
                   prompt_sha256=app.plain_v2_prompt_hash(self.settings.language))
        self.cp.data['identity'] = old
        self.cp.data['stages'] = {'0:translation': {'text': DRAFT, 'seconds': 40}}
        self.cp.data['pending'] = '0:fluency'
        self.cp.save()
        with contextlib.redirect_stdout(io.StringIO()):
            cp = app.Checkpoint(self.folder, self.identity)
        self.assertTrue((self.folder / 'checkpoint.before-fluency-update.json').exists())
        self.assertEqual(cp.data['pending'], '0:fluency')
        cp.data['pending'] = None  # Explicit retry consent.
        cp.save()
        text, fake = self.run_pipeline([PASS, PASS], cp)
        self.assertEqual(text, DRAFT + '\n')
        self.assertEqual(Counter(r for r, _ in fake.calls), Counter(fluency=1, accuracy=1))

    def test_fluency_uses_only_draft_accuracy_receives_source(self):
        _, fake = self.run_pipeline([DRAFT, PASS, PASS])
        by_role = {role: payload for role, payload in fake.calls}
        fluency = by_role['fluency']
        self.assertIn(DRAFT, fluency)
        self.assertNotIn('Mira arrived on Tuesday', fluency)
        self.assertNotIn('SOURCE:', fluency)
        self.assertIn('Mira arrived on Tuesday', by_role['accuracy'])
        self.assertIn(DRAFT, by_role['accuracy'])

    def test_truncated_fluency_records_exact_status_and_keeps_draft(self):
        with self.assertRaisesRegex(app.IncompleteResponse, '4096-token output budget'):
            self.run_pipeline({'translator': DRAFT,
                               'fluency': app.IncompleteResponse('fluency', 'length', 4096),
                               'accuracy': PASS})
        self.assertEqual(self.cp.data['pending'], '0:fluency')
        self.assertEqual(self.cp.data['last_response_failure']['completion_status'], 'length')
        self.assertEqual(set(self.cp.data['stages']), {'0:translation', '0:accuracy'})
        self.assertIn(DRAFT, (self.folder / 'translation.partial.txt').read_text(encoding='utf-8'))

    def test_v3_upgrade_resumes_each_auditor_stage_without_repeating_completed_calls(self):
        cases = (
            ('audit0', 'full', [DRAFT, PASS, PASS]),
            ('audit1', 'balanced', [DRAFT, FIX, PASS, FIXED]),
            ('audit2', 'balanced', [DRAFT, FIX, PASS, FIXED, FIX, FIXED]),
        )
        for stage, thoroughness, completed in cases:
            with self.subTest(stage=stage):
                self.settings.thoroughness = thoroughness
                identity = app.identity_for(self.source, self.chunks, self.configs, self.settings)
                folder = self.folder / stage
                folder.mkdir()
                cp = app.Checkpoint(folder, identity)
                with self.assertRaises(app.IncompleteResponse):
                    self.run_pipeline(completed + [app.IncompleteResponse('auditor', 'length', 2048)], cp)
                self.assertEqual(cp.data['pending'], '0:' + stage)
                self.assertNotIn('0:' + stage, cp.data['stages'])
                self.assertFalse((folder / 'translation.txt').exists())
                cp.data['identity'] = dict(identity, version='cicero-fluency-3',
                                           prompt_sha256=app.plain_v4_prompt_hash(self.settings.language))
                cp.save()
                before = cp.path.read_text(encoding='utf-8')
                completed_stages = dict(cp.data['stages'])
                with contextlib.redirect_stdout(io.StringIO()):
                    upgraded = app.Checkpoint(folder, identity)
                self.assertEqual(upgraded.data['stages'], completed_stages)
                self.assertEqual(upgraded.data['pending'], '0:' + stage)
                self.assertEqual((folder / 'checkpoint.before-review-update.json').read_text(encoding='utf-8'), before)
                upgraded.data['pending'] = None  # Explicit retry consent in the CLI.
                upgraded.save()
                result, fake = self.run_pipeline([PASS], upgraded)
                self.assertEqual([role for role, _ in fake.calls], ['auditor'])
                self.assertEqual(result, (DRAFT if stage == 'audit0' else FIXED) + '\n')
                self.assertTrue(upgraded.data['complete'])

    def test_v3_upgrade_preserves_completed_work_at_accuracy(self):
        with self.assertRaises(app.IncompleteResponse):
            self.run_pipeline([DRAFT, PASS, app.IncompleteResponse('accuracy', 'length', 2048)])
        self.cp.data['identity'] = dict(self.identity, version='cicero-fluency-3',
                                      prompt_sha256=app.plain_v4_prompt_hash(self.settings.language))
        self.cp.save()
        with contextlib.redirect_stdout(io.StringIO()):
            cp = app.Checkpoint(self.folder, self.identity)
        self.assertEqual(cp.data['pending'], '0:accuracy')
        cp.data['pending'] = None
        cp.save()
        text, fake = self.run_pipeline([PASS], cp)
        self.assertEqual(text, DRAFT + '\n')
        self.assertEqual([role for role, _ in fake.calls], ['accuracy'])

    def test_v3_upgrade_refuses_changed_prompts_or_configuration(self):
        previous = dict(self.identity, version='cicero-fluency-3',
                        prompt_sha256=app.plain_v4_prompt_hash(self.settings.language))
        for field, value in [('source_sha256', 'changed'), ('roles', {}),
                             ('chunks_sha256', 'changed'), ('prompt_sha256', 'unknown')]:
            with self.subTest(field=field):
                self.cp.data['identity'] = dict(previous, **{field: value})
                self.cp.save()
                before = self.cp.path.read_bytes()
                with self.assertRaises(app.Mismatch):
                    app.Checkpoint(self.folder, self.identity)
                self.assertEqual(self.cp.path.read_bytes(), before)


class FluencyResponseTests(unittest.TestCase):
    def make_models(self, response):
        from unittest.mock import Mock
        http = Mock()
        http.request.return_value = response
        return app.Models(http), http

    def cfg(self):
        return app.Config('google', 'catalog-selected-model', app.PROVIDERS['google'], SECRET, 128000, 4096)

    def envelope(self, text=PASS, finish='stop', refusal=None):
        return {'choices': [{'message': {'content': text, 'refusal': refusal}, 'finish_reason': finish}]}

    def test_all_roles_receive_their_configured_output_budget(self):
        for provider in ('google', 'openai', 'anthropic'):
            for allowance in (1024, 4096, 8192):
                with self.subTest(provider=provider, allowance=allowance):
                    cfg = app.Config(provider, 'catalog-selected-model', app.PROVIDERS[provider],
                                     SECRET, 128000, allowance)
                    response = ({'stop_reason': 'end_turn', 'content': [{'type': 'text', 'text': PASS}]}
                                if provider == 'anthropic' else self.envelope())
                    models, http = self.make_models(response)
                    for role in app.ROLES:
                        models.call(cfg, role, 'system', 'payload')
                    field = 'max_completion_tokens' if provider == 'openai' else 'max_tokens'
                    self.assertEqual(len(http.request.call_args_list), len(app.ROLES))
                    for request in http.request.call_args_list:
                        self.assertEqual(request.kwargs['json'][field], allowance)

    def test_google_connection_and_request_fields_are_preserved(self):
        cfg = self.cfg()
        models, http = self.make_models(self.envelope())
        models.call(cfg, 'fluency', 'system', 'draft')
        args, kwargs = http.request.call_args
        self.assertEqual(args, ('POST', cfg.base_url + '/chat/completions', {'Authorization': 'Bearer ' + SECRET}))
        self.assertEqual(kwargs['json'], {'model': cfg.model, 'stream': False, 'max_tokens': 4096,
            'messages': [{'role': 'system', 'content': 'system'}, {'role': 'user', 'content': 'draft'}]})

    def test_incomplete_verdict_is_not_accepted_even_if_pass_text_exists(self):
        for role in ('fluency', 'accuracy', 'auditor'):
            for finish, refusal, expected in [('length', None, 'length'), ('content_filter', None, 'content_filter'),
                                              ('stop', SECRET, 'refusal'), (SECRET, None, 'unknown')]:
                with self.subTest(role=role, expected=expected):
                    models, http = self.make_models(self.envelope(finish=finish, refusal=refusal))
                    with self.assertRaises(app.IncompleteResponse) as result:
                        models.call(self.cfg(), role, 'system', 'draft')
                    self.assertEqual(result.exception.reason, expected)
                    self.assertEqual(result.exception.role, role)
                    self.assertEqual(result.exception.output_budget, 4096)
                    self.assertNotIn(SECRET, str(result.exception))
                    http.request.assert_called_once()

    def test_complete_fluency_verdicts_pass_to_validation(self):
        for text in (PASS, FIX):
            models, _ = self.make_models(self.envelope(text=text))
            result = models.call(self.cfg(), 'fluency', app.system_prompt('fluency', 'Spanish'), DRAFT)
            self.assertEqual(app.validate_review(result)['ok'], text == PASS)

    def test_fluency_prompt_is_focused_and_keeps_style_context(self):
        prompt = app.system_prompt('fluency', 'Spanish')
        self.assertIn('Check ONLY the draft', prompt)
        self.assertIn('Do not assess source accuracy', prompt)
        self.assertIn('VERDICT: PASS', prompt)
        self.assertIn('VERDICT: FIX', prompt)
        payload = app.fluency_payload(DRAFT, 'A previous approved sentence.')
        self.assertIn('A previous approved sentence.', payload)
        self.assertIn(DRAFT, payload)
        self.assertNotIn('SOURCE:', payload)


class TextTests(unittest.TestCase):
    def test_review_convention_rejects_ambiguous_results(self):
        for bad in ('', 'PASS', 'VERDICT: PASS\nBut omissions remain.', 'VERDICT: FIX',
                    'VERDICT: FIX\nproblem\nVERDICT: PASS', '```\nVERDICT: PASS\n```'):
            with self.subTest(bad=bad), self.assertRaises(app.TranslationError):
                app.validate_review(bad)
        self.assertTrue(app.validate_review(PASS)['ok'])
        self.assertFalse(app.validate_review(FIX)['ok'])

    def test_wrappers_empty_text_and_extreme_content_loss_rejected(self):
        blocks = app.text_blocks(SOURCE)
        for bad in ('', '```\n' + DRAFT + '\n```', 'VERDICT: PASS\n' + DRAFT):
            with self.subTest(bad=bad), self.assertRaises(app.TranslationError):
                app.validate_translation(bad, blocks)
        with self.assertRaises(app.TranslationError):
            app.validate_translation('Resumen.', [app.Block('A complete source passage. ' * 100)])

    def test_paragraphs_markers_and_blank_lines_are_not_acceptance_rules(self):
        blocks = app.text_blocks(SOURCE)
        for draft in (DRAFT.replace('# Puerto', 'Puerto'), DRAFT.replace('- ', ''),
                      DRAFT.replace('\n\n', '\n'), DRAFT.replace('. ', '.\n\n'),
                      DRAFT.replace('\n\n', '\r\n\r\n')):
            with self.subTest(draft=draft):
                self.assertEqual(app.validate_translation(draft, blocks), draft.strip())

    def test_legacy_fingerprint_matches_the_archived_release(self):
        self.assertEqual(app.legacy_prompt_hash('Spanish'),
                         'dfd430f3ca875f1d86ef298b0b79c274516ad2465e1f65ec33594bf3a92d4abb')

    def test_long_paragraph_is_split_without_loss(self):
        for source in ('Long sentence. ' * 1000, '中文翻译文本' * 1000, 'слово ' * 1000):
            parts = list(app.split_block(app.Block(source), 350))
            self.assertEqual(''.join(b.text for b in parts), source)
            self.assertTrue(all(app.estimate_tokens(b.rendered()) <= 350 for b in parts))

    def test_target_is_tokens_not_8000_characters(self):
        cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], '', 128000, 24576)
        settings = app.Settings()
        self.assertEqual(app.source_limit(dict.fromkeys(app.ROLES, cfg), settings), 8000)
        chunks = app.make_chunks([app.Block('ordinary text ' * 1000)], 8000)
        self.assertEqual(len(chunks), 1)
        self.assertGreater(len(chunks[0][0].text), 8000)
        cfg.output_limit = 4096
        self.assertLess(app.source_limit(dict.fromkeys(app.ROLES, cfg), settings), 1600)

    def test_section_boundaries_and_internal_breaks(self):
        chunks = app.make_chunks([app.Block('one'), app.Block('chapter', 'heading', 1), app.Block('two')], 500)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(app.Block('one\n\ntwo').rendered(), 'one\ntwo')

    def test_encoding_bom_and_override(self):
        text = 'Привет мир'
        for encoding in ('utf-8', 'utf-8-sig', 'utf-16', 'utf-32'):
            self.assertEqual(app.decode_txt(text.encode(encoding))[0], text)
        self.assertEqual(app.decode_txt(text.encode('cp1251'), 'cp1251')[0], text)
        with self.assertRaises(UnicodeDecodeError):
            app.decode_txt(b'\xff', 'utf-8')

    def test_file_picker_manual_fallback(self):
        import sys
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.txt'
            path.write_text('sample')
            with patch.dict(sys.modules, {'tkinter': None}), patch.object(app, 'ask', return_value=str(path)):
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(app.select_document(), path.resolve())

    def test_cross_provider_destination_refused(self):
        cfg = app.Config('anthropic', 'test', app.PROVIDERS['google'], SECRET)
        with self.assertRaises(app.TranslationError):
            app.validate_destination(cfg)
        self.assertNotIn(SECRET, repr(cfg))

    def test_secret_in_public_settings_is_rejected(self):
        cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], SECRET)
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'source.txt'
            source.write_text('a source')
            with self.assertRaises(app.TranslationError):
                app.identity_for(source, [[app.Block('a source')]], {'translator': cfg}, app.Settings(language=SECRET))


class SetupTests(unittest.TestCase):
    def check_start_mode(self, mode):
        from unittest.mock import Mock
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / 'sample.txt'
            source.write_text(SOURCE, encoding='utf-8')
            cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], SECRET, 128000, 4096)
            configs = dict.fromkeys(app.ROLES, cfg)
            models = FakeModels([DRAFT, PASS, PASS])
            models.preflight = Mock()
            answers = ['no', 'Spanish', 'yes', mode] + (['none'] if mode != 'no' else [])
            with patch.object(app, '__file__', str(folder / 'interactive_translator.py')), \
                    patch.object(app, 'select_document', return_value=source), patch.object(app, 'HTTP', return_value=Mock()), \
                    patch.object(app, 'Models', return_value=models), patch.object(app, 'configure', return_value=configs), \
                    patch('builtins.input', side_effect=answers), contextlib.redirect_stdout(io.StringIO()):
                if mode == 'no':
                    with self.assertRaisesRegex(app.TranslationError, 'Stopped before inference'):
                        app.main()
                else:
                    app.main()
            if mode == 'yes':
                models.preflight.assert_called_once()
                self.assertEqual(models.preflight.call_args.args, (configs, 'Spanish'))
                self.assertIs(models.preflight.call_args.kwargs['retry_handler'], app.retry_failed_stage)
                self.assertTrue((models.preflight.call_args.kwargs['evidence_folder'] / 'checkpoint.json').exists())
            else:
                models.preflight.assert_not_called()
            outputs = list((folder / 'translations').glob('*/translation.txt'))
            if mode == 'no':
                self.assertEqual(models.calls, [])
                self.assertEqual(outputs, [])
            else:
                self.assertEqual(Counter(r for r, _ in models.calls), Counter(translator=1, fluency=1, accuracy=1))
                self.assertEqual(len(outputs), 1)
                self.assertEqual(outputs[0].read_text(encoding='utf-8'), DRAFT + '\n')

    def test_skip_checks_translates_and_keeps_normal_reviews(self):
        self.check_start_mode('skip')

    def test_yes_runs_checks_then_translates(self):
        self.check_start_mode('yes')

    def test_no_stops_without_checks_or_translation(self):
        self.check_start_mode('no')

    def test_same_model_asks_for_key_once(self):
        fake = type('Catalog', (), {'catalog': lambda self, cfg: ['catalog-model']})()
        with patch('builtins.input', side_effect=['yes', 'google', '', '1']), \
                patch.object(app, 'key_for', return_value=SECRET) as key, contextlib.redirect_stdout(io.StringIO()):
            configs = app.configure(fake, False)
        self.assertEqual(len(configs), 5)
        self.assertEqual(key.call_count, 1)
        self.assertEqual(len({id(c) for c in configs.values()}), 1)

    def test_individual_roles_reuse_provider_key(self):
        fake = type('Catalog', (), {'catalog': lambda self, cfg: ['catalog-model']})()
        with patch('builtins.input', side_effect=['no'] + ['google', '', '1'] * 3 +
                                               ['no'] + ['google', '', '1'] * 2), \
                patch.object(app, 'key_for', return_value=SECRET) as key, contextlib.redirect_stdout(io.StringIO()):
            configs = app.configure(fake, False)
        self.assertEqual(len(configs), 5)
        self.assertEqual(key.call_count, 1)

    def test_individual_role_descriptions_and_fixer_reuse(self):
        from unittest.mock import Mock
        fake = Mock()
        fake.catalog.side_effect = lambda cfg: [cfg.provider + '-selected-model']
        answers = ['no', '', '', '1', '', '', '1', '', '', '1', 'yes', '', '', '1']
        with patch('builtins.input', side_effect=answers), \
                patch.object(app, 'key_for', side_effect=lambda provider: SECRET + provider) as key, \
                contextlib.redirect_stdout(io.StringIO()) as output:
            configs = app.configure(fake, False)
        self.assertIs(configs['fixer'], configs['translator'])
        self.assertEqual([configs[r].provider for r in app.ROLES],
                         ['anthropic', 'google', 'deepseek', 'anthropic', 'openai'])
        self.assertEqual(key.call_count, 4)
        self.assertEqual(fake.catalog.call_count, 4)
        for label in ('Translator', 'Blind Critic', 'Alignment Inspector', 'Fixer', 'Final Auditor'):
            self.assertIn(label, output.getvalue())
        self.assertNotIn('Recommended:', output.getvalue())
        self.assertNotIn(SECRET, output.getvalue())

    def test_provider_menu_excludes_nvidia_and_defaults_to_google(self):
        fake = type('Catalog', (), {'catalog': lambda self, cfg: ['catalog-model']})()
        with patch('builtins.input', side_effect=['yes', 'nvidia', '', '', '1']) as inputs, \
                patch.object(app, 'key_for', return_value=SECRET) as key, \
                contextlib.redirect_stdout(io.StringIO()) as output:
            configs = app.configure(fake, False)
        self.assertNotIn('nvidia', app.PROVIDERS)
        provider_prompts = [call.args[0] for call in inputs.call_args_list if call.args[0].startswith('Provider ')]
        self.assertEqual(len(provider_prompts), 2)  # Removed option rejected; empty answer uses the default.
        self.assertTrue(all('nvidia' not in prompt and '[google]' in prompt for prompt in provider_prompts))
        self.assertEqual(configs['translator'].provider, 'google')
        key.assert_called_once_with('google')
        self.assertNotIn('Recommended:', output.getvalue())

    def test_advanced_document_settings_only_ask_about_relevant_file_types(self):
        cases = [('.txt', ['', '', '', '', '', 'cp1251']),
                 ('.PDF', ['', '', '', '', '', 'rus', 'yes']),
                 ('.docx', ['', '', '', '', '', 'eng+rus']),
                 ('.png', ['', '', '', '', '', 'eng+rus'])]
        for suffix, answers in cases:
            with self.subTest(suffix=suffix):
                settings = app.Settings()
                with patch('builtins.input', side_effect=answers) as inputs, contextlib.redirect_stdout(io.StringIO()):
                    timeout = app.advanced_document_settings(Path('source' + suffix), settings)
                prompts = '\n'.join(call.args[0] for call in inputs.call_args_list)
                self.assertEqual(timeout, 180)
                self.assertEqual(settings.source_target, 8000)
                self.assertEqual('TXT encoding override' in prompts, suffix == '.txt')
                self.assertEqual('Local OCR source-language' in prompts, suffix != '.txt')
                self.assertEqual('Force local OCR' in prompts, suffix == '.PDF')
                self.assertEqual(settings.force_ocr, suffix == '.PDF')
                if suffix == '.txt':
                    self.assertEqual(settings.encoding, 'cp1251')
                else:
                    self.assertEqual(settings.ocr_language, 'rus' if suffix == '.PDF' else 'eng+rus')

    def test_advanced_custom_model_keeps_user_selected_token_budgets(self):
        fake = type('Catalog', (), {'catalog': lambda self, cfg: []})()
        answers = ['yes', 'google', 'no', 'catalog-test-model', 'yes', '65536', '16000']
        with patch('builtins.input', side_effect=answers), patch.object(app, 'key_for', return_value=SECRET), \
                contextlib.redirect_stdout(io.StringIO()) as output:
            configs = app.configure(fake, True)
        for cfg in configs.values():
            self.assertEqual((cfg.context_limit, cfg.output_limit), (65536, 16000))
            self.assertEqual(cfg.model, 'catalog-test-model')
        self.assertNotIn(SECRET, output.getvalue())

    def test_custom_server_decline_stops_before_key_or_catalog(self):
        from unittest.mock import Mock
        fake = Mock()
        answers = ['yes', 'google', 'yes', 'https://example.invalid/v1', 'no']
        with patch('builtins.input', side_effect=answers), patch.object(app, 'key_for') as key, \
                contextlib.redirect_stdout(io.StringIO()), self.assertRaisesRegex(app.TranslationError, 'cancelled'):
            app.configure(fake, True)
        key.assert_not_called()
        fake.catalog.assert_not_called()

    def test_native_file_picker_and_cancel(self):
        import sys
        from types import SimpleNamespace
        from unittest.mock import Mock
        with tempfile.TemporaryDirectory() as folder:
            file = Path(folder) / 'source.txt'
            file.write_text('sample')
            root, dialog = Mock(), Mock(return_value=str(file))
            tk = SimpleNamespace(Tk=Mock(return_value=root), filedialog=SimpleNamespace(askopenfilename=dialog))
            with patch.dict(sys.modules, {'tkinter': tk, 'tkinter.filedialog': tk.filedialog}):
                self.assertEqual(app.select_document(), file.resolve())
                root.destroy.assert_called_once()
                dialog.return_value = ''
                with self.assertRaises(app.TranslationError):
                    app.select_document()

    def test_main_offline_end_to_end_and_completed_resume_without_key(self):
        import httpx
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / 'sample.txt'
            source.write_text(SOURCE, encoding='utf-8')
            sample = '# Puerto\n\nMira llegó el martes con 17 libros azules. No abrió la caja roja.'
            replies = iter([sample, PASS, PASS, FIX, sample, PASS, FIX, DRAFT, PASS, PASS])
            calls = []
            def handler(request):
                calls.append(request)
                if request.method == 'GET':
                    return httpx.Response(200, json={'data': [{'id': 'test-model'}]})
                return httpx.Response(200, json={'choices': [{'message': {'content': next(replies)}, 'finish_reason': 'stop'}],
                                                  'usage': {'prompt_tokens': 10, 'completion_tokens': 5}})
            http = app.HTTP(transport=httpx.MockTransport(handler))
            answers = ['no', 'Spanish', 'yes', 'google', '', '1', 'yes', 'yes', 'none']
            with patch.object(app, '__file__', str(folder / 'interactive_translator.py')), \
                    patch.object(app, 'select_document', return_value=source), patch.object(app, 'HTTP', return_value=http), \
                    patch.object(app, 'key_for', return_value=SECRET) as key, \
                    patch('builtins.input', side_effect=answers), contextlib.redirect_stdout(io.StringIO()):
                app.main()
                self.assertEqual(key.call_count, 1)
            outputs = list((folder / 'translations').glob('*/translation.txt'))
            self.assertEqual(len(outputs), 1)
            self.assertEqual(outputs[0].read_text(encoding='utf-8'), DRAFT + '\n')
            usage = json.loads((outputs[0].parent / 'usage.json').read_text(encoding='utf-8'))
            self.assertEqual(len(usage['requests']), 10)
            self.assertIn('Total tokens: 150', (outputs[0].parent / 'usage-summary.txt').read_text(encoding='utf-8'))
            total = len(calls)
            http = app.HTTP(transport=httpx.MockTransport(handler))
            with patch.object(app, '__file__', str(folder / 'interactive_translator.py')), \
                    patch.object(app, 'select_document', return_value=source), patch.object(app, 'HTTP', return_value=http), \
                    patch.object(app, 'key_for', side_effect=AssertionError('No key needed for complete run')), \
                    patch('builtins.input', side_effect=['1', '180', 'none']), contextlib.redirect_stdout(io.StringIO()):
                app.main()
            self.assertEqual(len(calls), total)


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_docx_export_and_paragraph_extraction_not_runs(self):
        from docx import Document
        source = self.folder / 'source.docx'
        doc = Document()
        doc.add_heading('Chapter one', level=1)
        p = doc.add_paragraph()
        p.add_run('The first ').bold = True
        p.add_run('paragraph has two runs.')
        doc.add_paragraph('First item', style='List Bullet')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'One table cell'
        doc.save(source)
        original = source.read_bytes()
        with contextlib.redirect_stdout(io.StringIO()):
            blocks = app.extract(source, app.Settings())
        self.assertEqual([b.text for b in blocks], ['Chapter one', 'The first paragraph has two runs.', 'First item', 'One table cell'])
        self.assertEqual(blocks[0].kind, 'heading')
        self.assertEqual(blocks[2].kind, 'list')
        destination = self.folder / 'translated.docx'
        app.export_docx(destination, DRAFT)
        result = Document(destination)
        self.assertEqual(result.paragraphs[0].style.name, 'Heading 1')
        self.assertEqual(result.paragraphs[1].text, app.text_blocks(DRAFT)[1].text)
        self.assertEqual(source.read_bytes(), original)

    def test_pdf_export_unicode_and_multi_page_content(self):
        import pymupdf as fitz
        source = '# Test chapter\n\n' + '\n\n'.join(f'Paragraph {i}: Café, mañana, Привет. ' + 'A long readable sentence. ' * 8 for i in range(45))
        destination = self.folder / 'translated.pdf'
        app.export_pdf(destination, source)
        with fitz.open(destination) as doc:
            self.assertGreater(len(doc), 1)
            text = ''.join(p.get_text() for p in doc)
            self.assertIn('Paragraph 0:', text)
            self.assertIn('Paragraph 44:', text)
            self.assertIn('Привет', text)

    def test_digital_and_scanned_pdf_are_not_duplicated(self):
        import pymupdf as fitz
        from PIL import Image
        source = self.folder / 'mixed.pdf'
        pixels = io.BytesIO()
        Image.new('RGB', (100, 100), 'white').save(pixels, format='PNG')
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 100), 'Digital paragraph with enough characters to avoid sparse classification.')
            page = doc.new_page()
            page.insert_image(page.rect, stream=pixels.getvalue())
            doc.save(source)
        original = source.read_bytes()
        with patch.object(app, 'local_ocr', return_value='Scanned paragraph') as ocr, contextlib.redirect_stdout(io.StringIO()):
            blocks = app.extract(source, app.Settings())
        text = '\n'.join(b.text for b in blocks)
        self.assertEqual(text.count('Digital paragraph'), 1)
        self.assertEqual(text.count('Scanned paragraph'), 1)
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual(source.read_bytes(), original)

    def test_image_and_image_only_docx_ocr_routes(self):
        from PIL import Image
        from docx import Document
        image = self.folder / 'scan.png'
        Image.new('RGB', (120, 90), 'white').save(image)
        source = self.folder / 'scan.docx'
        doc = Document()
        doc.add_picture(str(image))
        doc.save(source)
        with patch.object(app, 'local_ocr', return_value='A scanned paragraph') as ocr, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(app.extract(image, app.Settings())[0].text, 'A scanned paragraph')
            self.assertEqual(app.extract(source, app.Settings())[0].text, 'A scanned paragraph')
        self.assertEqual(ocr.call_count, 2)


class HTTPTests(unittest.TestCase):
    def setUp(self):
        import httpx
        self.httpx = httpx
        self.cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], SECRET, 128000, 4096)

    def with_handler(self, handler):
        http = app.HTTP(180, transport=self.httpx.MockTransport(handler))
        self.addCleanup(http.close)
        return app.Models(http)

    def envelope(self, text, finish='stop'):
        return self.httpx.Response(200, json={'choices': [{'message': {'content': text}, 'finish_reason': finish}]})

    def test_network_errors_not_retried_or_leaked(self):
        for cls in (self.httpx.ReadTimeout, self.httpx.ConnectTimeout, self.httpx.RemoteProtocolError):
            calls = []
            def handler(request):
                calls.append(request)
                raise cls(SECRET + ' DOCUMENT_CONTENT', request=request)
            models = self.with_handler(handler)
            with self.assertRaises(app.APIError) as exc:
                models.call(self.cfg, 'translator', 'system', 'DOCUMENT_CONTENT')
            self.assertEqual(len(calls), 1)
            self.assertIn(cls.__name__, str(exc.exception))
            self.assertNotIn(SECRET, str(exc.exception))
            self.assertNotIn('DOCUMENT_CONTENT', str(exc.exception))

    def test_404_410_and_rate_limit_not_replayed(self):
        for status in (404, 410, 429, 503, 202, 302):
            calls = []
            def handler(request):
                calls.append(request)
                return self.httpx.Response(status, json={'error': {'message': SECRET}})
            models = self.with_handler(handler)
            with self.assertRaises(app.APIError) as exc:
                models.call(self.cfg, 'translator', 'system', 'data')
            self.assertEqual(exc.exception.status, status)
            self.assertEqual(len(calls), 1)
            self.assertNotIn(SECRET, str(exc.exception))

    def test_200_empty_truncated_or_malformed_rejected(self):
        for response in (self.envelope(''), self.envelope(DRAFT, 'length'),
                         self.httpx.Response(200, content=b'not json'), self.httpx.Response(200, json={})):
            models = self.with_handler(lambda request: response)
            with self.assertRaises(app.TranslationError):
                models.call(self.cfg, 'translator', 'system', 'data')

    def test_configured_payload_and_provider_key_isolation(self):
        calls = []
        def handler(request):
            calls.append(request)
            return self.envelope(DRAFT)
        models = self.with_handler(handler)
        models.call(self.cfg, 'translator', 'system', 'plain source')
        body = json.loads(calls[0].content)
        self.assertEqual(body['max_tokens'], 4096)
        self.assertEqual(body['messages'][1]['content'], 'plain source')
        other = app.Config('glm', 'test-model', app.PROVIDERS['glm'], 'OTHER_FAKE_KEY', 128000, 24576)
        models.call(other, 'translator', 'system', 'plain source')
        self.assertEqual(calls[1].headers['authorization'], 'Bearer OTHER_FAKE_KEY')
        self.assertNotIn(SECRET, str(calls[1].headers))

    def test_openai_and_anthropic_adapters(self):
        calls = []
        def handler(request):
            calls.append(request)
            if request.url.path.endswith('/messages'):
                return self.httpx.Response(200, json={'stop_reason': 'end_turn', 'content': [{'type': 'text', 'text': DRAFT}]})
            return self.envelope(DRAFT)
        models = self.with_handler(handler)
        for provider in ('openai', 'anthropic'):
            base, model = app.PROVIDERS[provider], 'catalog-model'
            cfg = app.Config(provider, model, base, SECRET, 128000, 24576)
            self.assertEqual(models.call(cfg, 'translator', 'system', 'source'), DRAFT)
        self.assertIn('max_completion_tokens', json.loads(calls[0].content))
        self.assertEqual(calls[1].headers['x-api-key'], SECRET)
        self.assertNotIn('authorization', calls[1].headers)

    def test_preflight_checks_all_five_role_prompts(self):
        sample = '# Puerto\n\nMira llegó el martes con 17 libros azules. No abrió la caja roja.'
        replies = iter([sample, PASS, PASS, FIX, sample, PASS, FIX])
        calls = []
        def handler(request):
            calls.append(request)
            return self.envelope(next(replies))
        with contextlib.redirect_stdout(io.StringIO()):
            self.with_handler(handler).preflight(dict.fromkeys(app.ROLES, self.cfg), 'Spanish')
        self.assertEqual(len(calls), 7)

    def test_preflight_catalog_success_does_not_prove_inference(self):
        def handler(request):
            if request.method == 'GET':
                return self.httpx.Response(200, json={'data': [{'id': self.cfg.model}]})
            return self.httpx.Response(404, json={'error': {'message': 'Function not found for account'}})
        models = self.with_handler(handler)
        self.assertIn(self.cfg.model, models.catalog(self.cfg))
        with contextlib.redirect_stdout(io.StringIO()), self.assertRaises(app.APIError):
            models.preflight(dict.fromkeys(app.ROLES, self.cfg), 'Spanish')

    def test_preflight_rejects_reviewer_that_approves_everything(self):
        replies = iter(['# Puerto\n\nMira llegó el martes con 17 libros azules. No abrió la caja roja.', PASS, PASS, PASS])
        models = self.with_handler(lambda request: self.envelope(next(replies)))
        with contextlib.redirect_stdout(io.StringIO()), self.assertRaises(app.TranslationError):
            models.preflight(dict.fromkeys(app.ROLES, self.cfg), 'Spanish')

    def test_context_overflow_stops_before_http(self):
        calls = []
        models = self.with_handler(lambda request: calls.append(request))
        self.cfg.context_limit = 8192
        with self.assertRaises(app.TranslationError):
            models.call(self.cfg, 'translator', 'system', 'x' * 9000)
        self.assertFalse(calls)

    def test_timeout_configuration_reaches_client(self):
        http = app.HTTP(321, transport=self.httpx.MockTransport(lambda request: self.envelope(PASS)))
        self.addCleanup(http.close)
        self.assertEqual(http.client.timeout.read, 321)
        self.assertEqual(http.client.timeout.connect, 20)


if __name__ == '__main__':
    unittest.main()
