"""Regression tests for recovery, source-aware validation, startup roles and saved settings.

Run with: python -m pytest -q
All model replies are fake. No provider requests or keys are required.
"""
import json
import contextlib
import io
import threading
import time
from collections import Counter
from dataclasses import make_dataclass
from pathlib import Path
from unittest.mock import Mock, patch

import pytest

import interactive_translator as app
from test_translator import DRAFT, FIX, PASS, SECRET, SOURCE, FakeModels

SAMPLE = 'Puerto\n\nMira llegó el martes con 17 libros azules. No abrió la caja roja.'


@pytest.fixture
def saved_run(tmp_path):
    source = tmp_path / 'source.txt'
    source.write_text(SOURCE, encoding='utf-8')
    settings = app.Settings(language='Spanish')
    cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], SECRET, 128000, 4096)
    configs = dict.fromkeys(app.ROLES, cfg)
    chunks = [app.text_blocks(SOURCE)]
    cp = app.Checkpoint(tmp_path, app.identity_for(source, chunks, configs, settings))
    return cp, configs, settings, chunks


def test_reviewers_run_concurrently_within_a_passage(saved_run):
    cp, configs, settings, chunks = saved_run
    state = {'now': 0, 'peak': 0}
    lock = threading.Lock()

    def call(cfg, role, system, payload, **kwargs):
        if role == 'translator':
            return DRAFT
        with lock:
            state['now'] += 1
            state['peak'] = max(state['peak'], state['now'])
        time.sleep(0.05)  # Hold both reviewers open at once.
        with lock:
            state['now'] -= 1
        return PASS

    models = Mock()
    models.call.side_effect = call
    with contextlib.redirect_stdout(io.StringIO()):
        app.Pipeline(models, configs, settings, chunks, cp).run()
    assert state['peak'] == 2  # Blind Critic and Alignment Inspector overlapped.
    assert app.MAX_PARALLEL_CALLS == 10


@pytest.mark.parametrize('failure', [
    app.APIError(503, 'Service unavailable'),
    app.IncompleteResponse('fluency', 'length', 4096, 'VERDICT: FI'),
    app.IncompleteResponse('fluency', 'refusal', 4096),
    'This looks good but has no verdict.',
])
def test_retry_only_repeats_failed_stage_and_keeps_pending_until_success(saved_run, failure):
    cp, configs, settings, chunks = saved_run
    models = FakeModels({'translator': DRAFT, 'fluency': [failure, PASS], 'accuracy': PASS})
    def retry(label, error):
        assert label == 'Blind Critic'
        on_disk = json.loads(cp.path.read_text(encoding='utf-8'))
        # Both reviewers run together, so the other may still be in flight here.
        assert '0:fluency' in on_disk.get('pending_all', [])
        assert '0:translation' in on_disk['stages']
        assert not (cp.folder / 'translation.txt').exists()
        assert DRAFT in (cp.folder / 'translation.partial.txt').read_text(encoding='utf-8')
        return True
    handler = Mock(side_effect=retry)
    result = app.Pipeline(models, configs, settings, chunks, cp, retry_handler=handler).run()
    assert result == DRAFT + '\n'
    assert Counter(role for role, _ in models.calls) == Counter(translator=1, fluency=2, accuracy=1)
    handler.assert_called_once()
    assert cp.data['pending'] is None
    assert cp.data['complete']


def test_declined_retry_keeps_evidence_and_can_resume_later(saved_run):
    cp, configs, settings, chunks = saved_run
    models = FakeModels({'translator': DRAFT, 'fluency': 'Unstructured reviewer answer.', 'accuracy': PASS})
    with pytest.raises(app.ResponseValidationError):
        app.Pipeline(models, configs, settings, chunks, cp, retry_handler=lambda *_: False).run()
    assert cp.data['pending'] == '0:fluency'
    assert 'Unstructured reviewer answer.' in (cp.folder / 'passage-1-fluency.received.txt').read_text(encoding='utf-8')
    cp.data['pending'] = None  # Represents the user's explicit later retry choice.
    cp.save()
    models = FakeModels({'fluency': PASS})
    app.Pipeline(models, configs, settings, chunks, cp).run()
    assert [role for role, _ in models.calls] == ['fluency']
    archives = list(cp.folder.glob('passage-1-fluency-*.received.txt'))
    assert len(archives) == 2
    assert any('Unstructured reviewer answer.' in p.read_text(encoding='utf-8') for p in archives)


def test_later_passage_retry_does_not_repeat_previous_passage(saved_run):
    cp, configs, settings, _ = saved_run
    chunks = [[app.Block('First passage.')], [app.Block('Second passage.')]]
    source = cp.folder / 'source.txt'
    source.write_text('First passage.\n\nSecond passage.', encoding='utf-8')
    cp.data['identity'] = app.identity_for(source, chunks, configs, settings)
    cp.save()
    models = FakeModels({'translator': ['Primer pasaje.', 'Segundo pasaje.'],
                         'fluency': [PASS, 'Malformed review', PASS],
                         'accuracy': [PASS, PASS]})
    def retry(*_):
        assert '1:fluency' in cp.data.get('pending_all', [])
        assert app.Pipeline(models, configs, settings, chunks, cp).approved(0)
        return True
    result = app.Pipeline(models, configs, settings, chunks, cp, retry_handler=retry).run()
    assert result == 'Primer pasaje.\n\nSegundo pasaje.\n'
    assert Counter(role for role, _ in models.calls) == Counter(translator=2, fluency=3, accuracy=2)
    translations = [payload for role, payload in models.calls if role == 'translator']
    assert 'Primer pasaje.' in translations[1]


@pytest.mark.parametrize('received', [PASS + SECRET, app.IncompleteResponse('fluency', 'length', 4096, SECRET)])
def test_credentials_are_neither_saved_nor_offered_for_retry(saved_run, received):
    cp, configs, settings, chunks = saved_run
    handler = Mock(return_value=True)
    with pytest.raises(app.TranslationError, match='credential'):
        app.Pipeline(FakeModels({'translator': DRAFT, 'fluency': received, 'accuracy': PASS}),
                     configs, settings, chunks, cp, retry_handler=handler).run()
    handler.assert_not_called()
    assert all(SECRET not in file.read_text(encoding='utf-8') for file in cp.folder.iterdir() if file.is_file())


def test_local_evidence_write_failure_does_not_replay_a_paid_request(saved_run):
    cp, configs, settings, chunks = saved_run
    original = app.atomic_write
    def write(path, text):
        if str(path).endswith('.received.txt'):
            raise OSError('Simulated full disk')
        original(path, text)
    handler = Mock(return_value=True)
    models = FakeModels([DRAFT])
    with patch.object(app, 'atomic_write', side_effect=write), pytest.raises(OSError):
        app.Pipeline(models, configs, settings, chunks, cp, retry_handler=handler).run()
    handler.assert_not_called()
    assert len(models.calls) == 1
    assert cp.data['pending'] == '0:translation'


@pytest.mark.parametrize('provider', ['google', 'anthropic'])
def test_partial_response_text_is_retained_without_accepting_truncated_pass(provider):
    cfg = app.Config(provider, 'test-model', app.PROVIDERS[provider], SECRET, 128000, 4096)
    response = ({'stop_reason': 'max_tokens', 'content': [{'type': 'text', 'text': PASS}]}
                if provider == 'anthropic' else
                {'choices': [{'finish_reason': 'length', 'message': {'content': PASS}}]})
    http = Mock()
    http.request.return_value = response
    with pytest.raises(app.IncompleteResponse) as error:
        app.Models(http).call(cfg, 'auditor', 'system', 'draft')
    assert error.value.received_text == PASS
    assert PASS not in str(error.value)
    http.request.assert_called_once()


@pytest.mark.parametrize('text', [
    '```python\nprint("hello")\n```',
    '~~~text\nAn example\n~~~',
    'VERDICT: PASS\nThis is a literal label in the document.',
    '<think> is the tag described by this manual.',
    'Here is the translation is an example sentence.',
])
def test_literal_source_markers_are_not_rejected(text):
    assert app.validate_translation(text, [app.Block(text)]) == text


def test_literal_marker_can_move_to_a_new_line_without_becoming_a_wrapper():
    source = 'The source contains the label VERDICT: PASS.'
    translated = 'VERDICT: PASS\nThe label contained in the source.'
    assert app.validate_translation(translated, [app.Block(source)]) == translated


@pytest.mark.parametrize('text', ['```\nUna traducción.\n```', 'VERDICT: PASS', '<think>internal analysis'])
def test_markers_added_to_an_unrelated_source_are_still_rejected(text):
    with pytest.raises(app.ResponseValidationError):
        app.validate_translation(text, [app.Block('A simple source sentence.')])


def test_all_review_prompts_and_validation_share_one_findings_limit():
    for role in ('fluency', 'accuracy', 'auditor'):
        assert f'at most {app.REVIEW_FINDINGS_LIMIT} characters' in app.system_prompt(role, 'Spanish')
    assert not app.validate_review('VERDICT: FIX\n' + 'x' * app.REVIEW_FINDINGS_LIMIT)['ok']
    with pytest.raises(app.ResponseValidationError):
        app.validate_review('VERDICT: FIX\n' + 'x' * (app.REVIEW_FINDINGS_LIMIT + 1))
    assert not app.validate_review('VERDICT: FIX\nThe VERDICT: label from the source is missing.')['ok']
    with pytest.raises(app.ResponseValidationError):
        app.validate_review('VERDICT: FIX\nA problem.\nVERDICT: PASS')
    assert not app.validate_review('VERDICT: FIX\r\nA specific finding.')['ok']


@pytest.mark.parametrize('failed_role', ['fluency', 'fixer', 'auditor'])
def test_preflight_tests_assigned_roles_and_retries_only_failed_sample(tmp_path, failed_role):
    configs = {r: app.Config('google', 'test-' + r, app.PROVIDERS['google'], SECRET, 128000, 4096) for r in app.ROLES}
    models = app.Models(Mock())
    calls, successes = [], {r: 0 for r in app.ROLES}
    failed = False
    def call(cfg, role, system, payload, sample=False):
        nonlocal failed
        assert cfg is configs[role]
        assert system == app.system_prompt(role, 'Spanish')
        assert sample
        if role == 'fluency':
            assert 'SOURCE:' not in payload
        calls.append(role)
        if role == failed_role and not failed:
            failed = True
            return 'An invalid unstructured answer.' if role != 'fixer' else 'VERDICT: PASS'
        successes[role] += 1
        if role in ('translator', 'fixer'):
            return SAMPLE
        if role in ('accuracy', 'auditor'):
            return PASS if successes[role] == 1 else FIX
        return PASS
    retry = Mock(return_value=True)
    with patch.object(models, 'call', side_effect=call):
        models.preflight(configs, 'Spanish', retry_handler=retry, evidence_folder=tmp_path)
    retry.assert_called_once()
    assert calls.count('translator') == 1
    assert successes == {'translator': 1, 'fluency': 1, 'accuracy': 2, 'fixer': 1, 'auditor': 2}
    assert len(calls) == 8
    archives = list(tmp_path.glob('sample-*.received.txt'))
    assert any(('An invalid unstructured answer.' in p.read_text(encoding='utf-8') or
                (failed_role == 'fixer' and p.name.startswith('sample-fixer') and PASS in p.read_text(encoding='utf-8')))
               for p in archives)
    assert all(SECRET not in p.read_text(encoding='utf-8') for p in archives)


@pytest.mark.parametrize('answer, expected', [('yes', True), ('no', False), ('', False)])
def test_interactive_retry_requires_an_explicit_yes(monkeypatch, answer, expected):
    monkeypatch.setattr('builtins.input', lambda _: answer)
    assert app.retry_failed_stage('Blind Critic', app.APIError(503, 'Unavailable')) is expected


def test_saved_settings_use_defaults_for_missing_fields_and_reject_unknown_schema():
    settings = app.settings_from_saved({'language': 'Spanish'})
    assert settings.context_tokens == 1000 and settings.expansion == 2.5
    for saved in (None, [], {'future_field': SECRET}, {'context_tokens': '1000'},
                  {'expansion': float('nan')}, {'force_ocr': 'no'}):
        with pytest.raises(app.CheckpointSchemaError) as error:
            app.settings_from_saved(saved)
        assert SECRET not in str(error.value)


def test_a_future_required_setting_gets_a_schema_error(monkeypatch):
    monkeypatch.setattr(app, 'Settings', make_dataclass('FutureSettings', [('required_setting', str)]))
    with pytest.raises(app.CheckpointSchemaError, match='missing settings'):
        app.settings_from_saved({})


def test_resume_schema_error_happens_before_key_or_provider_request(saved_run, monkeypatch):
    cp, configs, settings, chunks = saved_run
    run_root = cp.folder / 'translations' / 'saved'
    run_root.mkdir(parents=True)
    cp.data['identity']['settings']['future_setting'] = 'unknown'
    checkpoint = run_root / 'checkpoint.json'
    checkpoint.write_text(json.dumps(cp.data), encoding='utf-8')
    original = checkpoint.read_bytes()
    monkeypatch.setattr(app, '__file__', str(cp.folder / 'interactive_translator.py'))
    monkeypatch.setattr(app, 'select_document', lambda: cp.folder / 'source.txt')
    monkeypatch.setattr('builtins.input', lambda _: '1')
    with patch.object(app, 'HTTP') as http, patch.object(app, 'key_for') as key, pytest.raises(app.CheckpointSchemaError):
        app.main()
    http.assert_not_called()
    key.assert_not_called()
    assert checkpoint.read_bytes() == original


def test_saved_model_schema_reports_role_and_type_problems(saved_run):
    cp, configs, _, _ = saved_run
    saved = {r: cfg.public() for r, cfg in configs.items()}
    assert set(app.configs_from_saved(saved)) == set(app.ROLES)
    for invalid in ({}, dict(saved, extra={}), dict(saved, translator={'provider': 'google'}),
                    dict(saved, translator=dict(saved['translator'], output_limit='4096'))):
        with pytest.raises(app.CheckpointSchemaError):
            app.configs_from_saved(invalid)


@pytest.mark.parametrize('answer', ['cancel', 'CANCEL', 'q'])
def test_numeric_prompts_can_cancel_without_ctrl_c(monkeypatch, answer):
    monkeypatch.setattr('builtins.input', lambda _: answer)
    with pytest.raises(app.TranslationError, match='Cancelled'):
        app.number('Passage size', 8000, 100, 8000)


def test_expansion_and_continuity_controls_affect_capacity(monkeypatch):
    answers = iter(['180', 'balanced', '8000', '3.5', '4000', 'auto'])
    monkeypatch.setattr('builtins.input', lambda _: next(answers))
    settings = app.Settings()
    app.advanced_document_settings(Path('source.txt'), settings)
    assert settings.expansion == 3.5 and settings.context_tokens == 4000
    cfg = app.Config('google', 'test-model', app.PROVIDERS['google'], context_limit=32768, output_limit=4096)
    configs = dict.fromkeys(app.ROLES, cfg)
    assert app.source_limit(configs, settings) < app.source_limit(configs, app.Settings())
    settings.expansion = 2.5
    assert app.source_limit(configs, settings) < app.source_limit(configs, app.Settings())


def test_invalid_decimal_input_reprompts(monkeypatch):
    answers = iter(['nan', 'inf', '0', '2.75'])
    monkeypatch.setattr('builtins.input', lambda _: next(answers))
    assert app.number('Expansion', 2.5, 1, 8, decimal=True) == 2.75


def test_extremely_large_integer_input_reprompts_without_overflow(monkeypatch):
    answers = iter(['9' * 1000, '180'])
    monkeypatch.setattr('builtins.input', lambda _: next(answers))
    assert app.number('Timeout', 180, 10, 3600) == 180


def test_main_forwards_user_timeout_to_extraction(saved_run, monkeypatch):
    cp, configs, _, chunks = saved_run
    source = cp.folder / 'source.pdf'
    source.write_bytes(b'A fake PDF, read only through the extraction mock.')
    monkeypatch.setattr(app, '__file__', str(cp.folder / 'interactive_translator.py'))
    monkeypatch.setattr(app, 'select_document', lambda: source)
    answers = iter(['yes', 'Spanish', '321', 'balanced', '8000', '2.5', '1000', 'eng', 'no', 'skip', 'none'])
    monkeypatch.setattr('builtins.input', lambda _: next(answers))
    with patch.object(app, 'HTTP', return_value=Mock()) as http, \
            patch.object(app, 'Models', return_value=FakeModels([DRAFT, PASS, PASS])), \
            patch.object(app, 'configure', return_value=configs), patch.object(app, 'extract', return_value=chunks[0]) as extract:
        app.main()
    http.assert_called_once_with(321)
    assert extract.call_args.kwargs['ocr_timeout'] == 321


def test_ocr_uses_configured_timeout():
    with patch('pytesseract.image_to_string', return_value='Recognized text') as ocr:
        assert app.local_ocr(object(), 'rus', timeout=37) == 'Recognized text'
    assert ocr.call_args.kwargs == {'lang': 'rus', 'timeout': 37}


@pytest.mark.parametrize('suffix', ['.png', '.docx', '.pdf'])
def test_ocr_timeout_is_forwarded_by_each_extraction_route(tmp_path, suffix):
    from PIL import Image
    from docx import Document
    import pymupdf
    image = tmp_path / 'scan.png'
    Image.new('RGB', (80, 80), 'white').save(image)
    source = tmp_path / ('source' + suffix)
    if suffix == '.png':
        source = image
    elif suffix == '.docx':
        doc = Document()
        doc.add_picture(str(image))
        doc.save(source)
    else:
        with pymupdf.open() as doc:
            doc.new_page().insert_image(pymupdf.Rect(0, 0, 100, 100), filename=str(image))
            doc.save(source)
    with patch.object(app, 'local_ocr', return_value='Recognized source text') as ocr:
        blocks = app.extract(source, app.Settings(), ocr_timeout=37)
    assert blocks[0].text == 'Recognized source text'
    assert ocr.call_args.kwargs['timeout'] == 37


def test_numbered_lists_do_not_gain_bullets_in_rendering_or_exports(tmp_path):
    from docx import Document
    import pymupdf
    text = '1. First item\n2. Second item\n\n- A bullet item'
    blocks = app.text_blocks(text)
    assert blocks[0].rendered() == '1. First item'
    assert blocks[0].budget_text() == '- 1. First item'  # Historical conservative chunk accounting only.
    assert blocks[2].rendered() == '- A bullet item'
    word = tmp_path / 'numbered.docx'
    pdf = tmp_path / 'numbered.pdf'
    app.export_docx(word, text)
    doc = Document(word)
    assert doc.paragraphs[0].text == '1. First item'
    assert doc.paragraphs[0].style.name == 'Normal'
    assert doc.paragraphs[2].style.name == 'List Bullet'
    app.export_pdf(pdf, text)
    with pymupdf.open(pdf) as doc:
        rendered = ''.join(page.get_text() for page in doc)
    assert '1. First item' in rendered
    assert rendered.count('•') == 1


@pytest.mark.parametrize('text', ['a' * 9000, 'ordinary words ' * 1000, '中文文本' * 1000, 'слово ' * 1000])
def test_chunk_split_covers_every_character(text):
    parts = list(app.split_block(app.Block(text), 350))
    assert ''.join(part.text for part in parts) == text
    assert all(app.estimate_tokens(part.budget_text()) <= 350 for part in parts)


def test_uncertain_legacy_encoding_is_reported_instead_of_guessed():
    with patch('charset_normalizer.from_bytes') as detect:
        detect.return_value.best.return_value = None
        with pytest.raises(app.TranslationError, match='uncertain'):
            app.decode_txt(b'\x81\xff\xab')


@pytest.mark.parametrize('raw', [b'\x81\xff\xab', b'hello\x00'])
def test_auto_decoding_uses_a_confident_legacy_detection(raw):
    class DetectedText:
        chaos = 0.01
        encoding = 'cp1251'
        def __str__(self):
            return 'Привет'
    with patch('charset_normalizer.from_bytes') as detect:
        detect.return_value.best.return_value = DetectedText()
        assert app.decode_txt(raw) == ('Привет', 'cp1251')
        detect.assert_called_once_with(raw)


@pytest.mark.parametrize('value', [None, [], 17])
def test_non_text_review_cannot_be_approval(value):
    with pytest.raises(app.ResponseValidationError):
        app.validate_review(value)


def test_empty_chunk_input_has_no_phantom_passage():
    assert app.make_chunks([], 350) == []


def test_startup_number_check_accepts_localized_digits(saved_run):
    cp, configs, _, _ = saved_run
    models = app.Models(Mock())
    replies = iter([SAMPLE.replace('17', '١٧'), PASS, PASS, FIX, SAMPLE, PASS, FIX])
    with patch.object(models, 'call', side_effect=lambda *a, **kw: next(replies)) as call:
        models.preflight(configs, 'Spanish', evidence_folder=cp.folder)
    assert '999' in call.call_args_list[3].args[3]


def test_v4_checkpoint_upgrade_keeps_approved_work_and_rejects_other_changes(saved_run):
    cp, configs, settings, chunks = saved_run
    current = cp.data['identity']
    previous = dict(current, version='cicero-review-4', prompt_sha256=app.plain_v4_prompt_hash(settings.language))
    assert app.compatible_text_only_upgrade(previous, current)
    for key in ('source_sha256', 'chunks_sha256', 'prompt_sha256'):
        assert not app.compatible_text_only_upgrade(dict(previous, **{key: 'changed'}), current)
    cp.data['identity'] = previous
    cp.data['stages'] = {'0:translation': {'text': DRAFT, 'seconds': 1}}
    cp.data['pending'] = '0:fluency'
    cp.save()
    original = cp.path.read_text(encoding='utf-8')
    upgraded = app.Checkpoint(cp.folder, current)
    assert upgraded.data['pending'] == '0:fluency'
    assert upgraded.data['stages']['0:translation']['text'] == DRAFT
    assert (cp.folder / 'checkpoint.before-reliability-update.json').read_text(encoding='utf-8') == original
